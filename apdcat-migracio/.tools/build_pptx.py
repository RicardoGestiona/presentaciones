"""
build_pptx.py — PPTX nativo y editable de apdcat-migracio, fiel al render HTML.

Fuente: index-standalone.html (offline, mismas fuentes e iconos que el deck).
Un walker DOM extrae cada elemento visible como átomo (forma, texto o icono) con
la posición y el estilo calculados por el navegador; python-pptx los reconstruye
uno a uno, así que cada elemento es seleccionable y editable en PowerPoint.

Fidelidad:
  · Diapositiva de 1280×720 pt → 1 px CSS = 1 pt (regla CLAUDE.md §8: 1rem = 16pt).
  · Los pseudo-elementos ::before/::after se materializan como nodos reales antes
    de extraer (puntos y línea de la hoja de ruta, círculos decorativos…).
  · Iconos Font Awesome → PNG transparente a 4× (PowerPoint no tiene la fuente).
  · Degradados radiales, sombras (incluidos anillos 0 0 0 Npx), opacidad,
    rotación (rombos del gantt) y colores color-mix() / color(srgb …).
Animaciones:
  · Cada .phase[data-phase] se agrupa y recibe la entrada nativa «Flotar hacia
    dentro» (fundido + subida) al clic; los del mismo data-phase, a la vez.
  · Transición «Empujar» entre diapositivas (equivale al desplazamiento horizontal).
  · No hay clases animate-* en el deck (regla §8.2 sin aplicación).

Salida: apdcat-migracio/apdcat-migracio.pptx

Uso (reutiliza el venv de apdcat/.tools, con python-pptx y playwright; navegador: Google Chrome):
    apdcat/.tools/.venv/bin/python -I apdcat-migracio/.tools/build_standalone.py
    apdcat/.tools/.venv/bin/python -I apdcat-migracio/.tools/build_pptx.py
"""
from __future__ import annotations

import json
import math
import sys
from io import BytesIO
from pathlib import Path

from lxml import etree
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT     = Path(__file__).resolve().parent.parent
SOURCE   = ROOT / "index-standalone.html"
OUT_PPTX = ROOT / "apdcat-migracio.pptx"

W_CSS, H_CSS = 1280, 720
EMU_PER_PX   = 12700            # 1 px CSS = 1 pt
FONT         = "Roboto"
ICON_SCALE   = 4

PHASE_DURATION_MS      = 500    # = transición .phase de corporate.css
PHASE_RISE             = 12 / H_CSS   # translateY(12px) en fracción de alto de diapositiva
SLIDE_TRANSITION_MS    = 700    # = transición de main[role=presentation]

NS_P   = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A   = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_MC  = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def log(event: str, **fields) -> None:
    print(json.dumps({"event": event, **fields}, ensure_ascii=False), file=sys.stderr)


def px(n: float) -> Emu:
    return Emu(int(round(n * EMU_PER_PX)))


# ─── JS: estado final + pseudo-elementos materializados ───────────────────
JS_PREPARE = r"""
() => {
    const style = document.createElement('style');
    style.textContent = `*{transition:none!important;animation:none!important}
        .__np::before,.__np::after{content:none!important}`;
    document.head.appendChild(style);
    document.querySelectorAll('.phase').forEach(el => el.classList.add('visible'));

    // Copia cada ::before/::after a un <span> real con su estilo calculado, para que
    // el walker obtenga su caja con getBoundingClientRect().
    const targets = [];
    document.querySelectorAll('.slide-container, .slide-container *').forEach(el => {
        for (const which of ['before', 'after']) {
            const st = getComputedStyle(el, '::' + which);
            const content = st.content;
            if (!content || content === 'none' || content === 'normal') continue;
            targets.push([el, which, st]);
        }
    });
    for (const [el, which, st] of targets) {
        const span = document.createElement('span');
        for (const prop of st) span.style.setProperty(prop, st.getPropertyValue(prop));
        const text = st.content.replace(/^["']|["']$/g, '')
            .replace(/\\([0-9a-fA-F]+)\s?/g, (m, hex) => String.fromCodePoint(parseInt(hex, 16)));
        span.textContent = text;
        span.dataset.pseudo = which;
        if (which === 'before') el.prepend(span); else el.append(span);
        el.classList.add('__np');
    }
    return targets.length;
}
"""

# ─── JS: extracción de átomos de una diapositiva ──────────────────────────
JS_EXTRACT = r"""
(slideIndex) => {
    const slide = document.querySelectorAll('.slide-container')[slideIndex];
    const sr = slide.getBoundingClientRect();
    const phaseEls = Array.from(slide.querySelectorAll('.phase[data-phase]'));
    const atoms = [], backAtoms = [];
    const SKIP_TAGS = new Set(['SCRIPT', 'STYLE', 'LINK', 'META', 'BR']);

    // rgb()/rgba() o color(srgb r g b / a) → {c:[r,g,b], a}
    function color(s) {
        if (!s) return null;
        let m = s.match(/^rgba?\(([^)]+)\)/);
        if (m) {
            const p = m[1].split(/[\s,\/]+/).filter(Boolean).map(parseFloat);
            return { c: [p[0], p[1], p[2]].map(Math.round), a: p.length > 3 ? p[3] : 1 };
        }
        m = s.match(/^color\(srgb\s+([^)]+)\)/);
        if (m) {
            const p = m[1].split(/[\s\/]+/).filter(Boolean).map(parseFloat);
            return { c: [p[0], p[1], p[2]].map(v => Math.round(v * 255)), a: p.length > 3 ? p[3] : 1 };
        }
        return null;
    }
    function gradientStops(bgImage) {
        if (!bgImage || !bgImage.startsWith('radial-gradient')) return null;
        const stops = bgImage.match(/(rgba?\([^)]*\)|color\(srgb[^)]*\))/g) || [];
        return stops.map(color).filter(Boolean);
    }
    function shadow(s) {
        if (!s || s === 'none' || s.includes('inset')) return null;
        const m = s.match(/(rgba?\([^)]*\)|color\(srgb[^)]*\))\s+(-?[\d.]+)px\s+(-?[\d.]+)px\s+([\d.]+)px(?:\s+(-?[\d.]+)px)?/);
        if (!m) return null;
        return { color: color(m[1]), x: +m[2], y: +m[3], blur: +m[4], spread: m[5] ? +m[5] : 0 };
    }
    function rotationDeg(t) {
        const m = t && t.match(/^matrix\(([^)]+)\)/);
        if (!m) return 0;
        const [a, b] = m[1].split(',').map(parseFloat);
        return Math.round(Math.atan2(b, a) * 180 / Math.PI * 100) / 100;
    }
    function phaseOf(el) {
        const ph = el.closest('.phase[data-phase]');
        if (!ph || !slide.contains(ph)) return null;
        return { key: phaseEls.indexOf(ph), num: +ph.dataset.phase };
    }
    function box(el, st) {
        const r = el.getBoundingClientRect();
        const rot = rotationDeg(st.transform);
        if (rot) {  // caja sin rotar centrada en la caja rotada
            const w = el.offsetWidth, h = el.offsetHeight;
            return { x: r.left - sr.left + (r.width - w) / 2, y: r.top - sr.top + (r.height - h) / 2, w, h, rot };
        }
        return { x: r.left - sr.left, y: r.top - sr.top, w: r.width, h: r.height, rot: 0 };
    }

    function textStyle(st) {
        let lh = parseFloat(st.lineHeight);
        if (isNaN(lh)) lh = parseFloat(st.fontSize) * 1.2;
        return {
            fontSize: parseFloat(st.fontSize),
            color: (color(st.color) || { c: [0, 0, 0] }).c,
            align: st.textAlign,
            lineHeight: lh,
            uppercase: st.textTransform === 'uppercase',
            letterSpacing: parseFloat(st.letterSpacing) || 0,
            pad: [st.paddingTop, st.paddingRight, st.paddingBottom, st.paddingLeft].map(v => parseFloat(v) || 0),
            justify: st.justifyContent, alignItems: st.alignItems, display: st.display,
        };
    }
    function runs(el) {
        const out = [];
        (function walk(node, ctx) {
            if (node.nodeType === 3) {
                if (node.textContent.length) out.push({ text: node.textContent, ...ctx });
                return;
            }
            if (node.nodeType !== 1 || node.tagName === 'I') return;
            if (node.tagName === 'BR') { out.push({ text: '\n', ...ctx }); return; }
            const st = getComputedStyle(node);
            if (st.display === 'none') return;
            const next = {
                bold: (parseInt(st.fontWeight) || 400) >= 600,
                italic: st.fontStyle === 'italic',
                color: (color(st.color) || { c: ctx.color }).c,
                size: parseFloat(st.fontSize),
                spacing: parseFloat(st.letterSpacing) || 0,
                upper: st.textTransform === 'uppercase',
            };
            for (const c of node.childNodes) walk(c, next);
        })(el, null);
        return out;
    }

    function walk(el, opacity) {
        if (SKIP_TAGS.has(el.tagName)) return;
        const st = getComputedStyle(el);
        if (st.display === 'none' || st.visibility === 'hidden') return;
        const op = opacity * (parseFloat(st.opacity) || 0);
        if (op < 0.005) return;
        const phase = phaseOf(el);

        // Icono Font Awesome → imagen.
        if (el.tagName === 'I' && /\bfa-/.test(el.className)) {
            const r = el.getBoundingClientRect();
            atoms.push({ kind: 'icon', x: r.left - sr.left, y: r.top - sr.top, w: r.width, h: r.height,
                         cls: el.className, color: st.color, fontSize: st.fontSize, phase });
            return;
        }

        const b = box(el, st);
        const isSlide = el === slide;
        if (b.w >= 0.5 && b.h >= 0.5 && !isSlide) {
            const bg = color(st.backgroundColor);
            const grad = gradientStops(st.backgroundImage);
            const borders = ['Top', 'Right', 'Bottom', 'Left'].map(s => ({
                w: parseFloat(st['border' + s + 'Width']) || 0,
                c: color(st['border' + s + 'Color']),
                style: st['border' + s + 'Style'],
            })).map(x => (x.style === 'none' || !x.c || x.c.a < 0.01) ? { w: 0, c: null } : x);
            const hasBg = (bg && bg.a > 0.01) || (grad && grad.length > 1);
            const hasBorder = borders.some(x => x.w > 0.5);
            if (hasBg || hasBorder) {
                const atom = { kind: 'shape', ...b, bg: hasBg && !grad ? bg : null, grad, opacity: op,
                               borders, radius: parseFloat(st.borderTopLeftRadius) || 0,
                               shadow: shadow(st.boxShadow), phase };
                // Decoración de fondo de la diapositiva (z-index 0): debajo de todo.
                if (el.dataset.pseudo && el.parentElement === slide) backAtoms.push(atom); else atoms.push(atom);
            }
        }

        const kids = Array.from(el.children).filter(c => getComputedStyle(c).display !== 'none');
        const blockKids = kids.filter(c => {
            if (c.tagName === 'I') return true;
            const d = getComputedStyle(c).display;
            return !['inline', 'inline-block', 'inline-flex', 'contents'].includes(d);
        });
        const flexLike = ['flex', 'inline-flex', 'grid', 'inline-grid'].includes(st.display);

        if (blockKids.length === 0 && !flexLike && el.innerText && el.innerText.trim()) {
            atoms.push({ kind: 'text', ...b, ...textStyle(st), runs: runs(el), opacity: op, phase });
            return;
        }
        // Nodos de texto sueltos junto a hijos de bloque (p. ej. icono + texto en flex).
        for (const node of el.childNodes) {
            if (node.nodeType !== 3 || !node.textContent.trim()) continue;
            const range = document.createRange(); range.selectNodeContents(node);
            const r = range.getBoundingClientRect();
            const ts = textStyle(st);
            atoms.push({ kind: 'text', x: r.left - sr.left, y: r.top - sr.top, w: r.width + 2, h: r.height, rot: 0,
                         ...ts, pad: [0, 0, 0, 0], align: 'left',
                         runs: [{ text: node.textContent.trim(), bold: (parseInt(st.fontWeight) || 400) >= 600,
                                  italic: false, color: ts.color, size: ts.fontSize, spacing: ts.letterSpacing,
                                  upper: ts.uppercase }],
                         opacity: op, phase });
        }
        for (const c of kids) walk(c, op);
    }

    walk(slide, 1);
    const slideBg = color(getComputedStyle(slide).backgroundColor);
    return { atoms: backAtoms.concat(atoms), bg: slideBg ? slideBg.c : [255, 255, 255],
             phases: phaseEls.map(p => +p.dataset.phase) };
}
"""

ICON_PAGE = """<!doctype html><html><head><meta charset="utf-8">{fa_css}
<style>html,body{{margin:0;background:transparent}} i{{display:inline-block;line-height:1}}</style>
</head><body><i id="icon"></i></body></html>"""


# ─── helpers DrawingML ────────────────────────────────────────────────────

def _hex(c) -> str:
    return f"{int(c[0]):02X}{int(c[1]):02X}{int(c[2]):02X}"


def _sppr(shape):
    return shape._element.find(qn("p:spPr"))


def _clear_fill(spPr) -> None:
    for tag in ("a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill", "a:grpFill", "a:noFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)


def _insert_fill(spPr, fill_el) -> None:
    """La relleno va justo tras la geometría (orden del esquema CT_ShapeProperties)."""
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    geom.addnext(fill_el)


def _srgb(parent, c, alpha: float = 1.0):
    el = etree.SubElement(parent, qn("a:srgbClr"), val=_hex(c))
    if alpha < 0.999:
        etree.SubElement(el, qn("a:alpha"), val=str(int(round(alpha * 100000))))
    return el


def set_solid_fill(shape, c, alpha: float = 1.0) -> None:
    spPr = _sppr(shape)
    _clear_fill(spPr)
    fill = etree.Element(qn("a:solidFill"))
    _srgb(fill, c, alpha)
    _insert_fill(spPr, fill)


def set_radial_fill(shape, stops, alpha: float) -> None:
    spPr = _sppr(shape)
    _clear_fill(spPr)
    grad = etree.Element(qn("a:gradFill"), rotWithShape="1")
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for i, stop in enumerate(stops):
        pos = int(round(i / (len(stops) - 1) * 100000))
        gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(pos))
        _srgb(gs, stop["c"], alpha * stop["a"])
    path = etree.SubElement(grad, qn("a:path"), path="circle")
    etree.SubElement(path, qn("a:fillToRect"), l="50000", t="50000", r="50000", b="50000")
    _insert_fill(spPr, grad)


def set_no_fill(shape) -> None:
    spPr = _sppr(shape)
    _clear_fill(spPr)
    _insert_fill(spPr, etree.Element(qn("a:noFill")))


def set_line(shape, c=None, width_pt: float = 0.0, alpha: float = 1.0) -> None:
    spPr = _sppr(shape)
    ln = spPr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(spPr, qn("a:ln"))
    for ch in list(ln):
        ln.remove(ch)
    if c is None:
        etree.SubElement(ln, qn("a:noFill"))
        return
    ln.set("w", str(int(width_pt * 12700)))
    fill = etree.SubElement(ln, qn("a:solidFill"))
    _srgb(fill, c, alpha)


def set_outer_shadow(shape, sh: dict, alpha: float) -> None:
    spPr = _sppr(shape)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    dist = math.hypot(sh["x"], sh["y"])
    direction = (math.degrees(math.atan2(sh["y"], sh["x"])) % 360) if dist else 0
    shdw = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(sh["blur"] * 12700)),
                            dist=str(int(dist * 12700)), dir=str(int(direction * 60000)),
                            algn="ctr", rotWithShape="0")
    _srgb(shdw, sh["color"]["c"], sh["color"]["a"] * alpha)


def _rect_shape(shapes, x, y, w, h, radius: float = 0.0):
    m = min(w, h)
    if radius > 0 and m > 0 and radius >= m / 2 - 0.5 and abs(w - h) < 1:
        s = shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(y), px(w), px(h))
    elif radius > 0 and m > 0:
        s = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h))
        s.adjustments[0] = min(radius / m, 0.5)
    else:
        s = shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    s.shadow.inherit = False
    return s


def _strip_text_frame(shape) -> None:
    shape.text_frame.text = ""


# ─── átomos → formas ──────────────────────────────────────────────────────

def add_shape_atom(shapes, a: dict) -> None:
    op = a["opacity"]
    sh = a.get("shadow")
    # Anillo (0 0 0 Npx): forma ampliada detrás, del color de la sombra.
    if sh and sh["blur"] == 0 and sh["x"] == 0 and sh["y"] == 0 and sh["spread"] > 0:
        sp = sh["spread"]
        ring = _rect_shape(shapes, a["x"] - sp, a["y"] - sp, a["w"] + 2 * sp, a["h"] + 2 * sp,
                           a["radius"] + sp if a["radius"] else 0)
        set_solid_fill(ring, sh["color"]["c"], sh["color"]["a"] * op)
        set_line(ring)
        _strip_text_frame(ring)
        sh = None

    borders = a["borders"]
    top, right, bottom, left = borders
    uniform = (all(b["w"] > 0.5 for b in borders)
               and len({(round(b["w"], 1), tuple(b["c"]["c"])) for b in borders}) == 1)
    only_top = top["w"] > 0.5 and not any(b["w"] > 0.5 for b in (right, bottom, left))

    if only_top and a["radius"] > 0:
        # Tarjeta redondeada con borde superior: base del color del borde y cuerpo encima
        # desplazado su grosor → el borde se curva en las esquinas como en CSS.
        base = _rect_shape(shapes, a["x"], a["y"], a["w"], a["h"], a["radius"])
        set_solid_fill(base, top["c"]["c"], top["c"]["a"] * op)
        set_line(base)
        if sh and sh["color"]:
            set_outer_shadow(base, sh, op)
        _strip_text_frame(base)
        body = _rect_shape(shapes, a["x"], a["y"] + top["w"], a["w"], a["h"] - top["w"], a["radius"])
        if a["bg"]:
            set_solid_fill(body, a["bg"]["c"], a["bg"]["a"] * op)
        else:
            set_solid_fill(body, [255, 255, 255], op)
        set_line(body)
        _strip_text_frame(body)
        return

    s = _rect_shape(shapes, a["x"], a["y"], a["w"], a["h"], a["radius"])
    if a.get("rot"):
        s.rotation = a["rot"]
    if a["grad"] and len(a["grad"]) > 1:
        set_radial_fill(s, a["grad"], op)
    elif a["bg"]:
        set_solid_fill(s, a["bg"]["c"], a["bg"]["a"] * op)
    else:
        set_no_fill(s)
    if uniform:
        b = borders[0]
        set_line(s, b["c"]["c"], b["w"], b["c"]["a"] * op)
    else:
        set_line(s)
    if sh and sh["color"]:
        set_outer_shadow(s, sh, op)
    _strip_text_frame(s)
    if uniform:
        return

    # Bordes por lado: barras independientes (filetes, barra del título…).
    x, y, w, h = a["x"], a["y"], a["w"], a["h"]
    for b, (bx, by, bw, bh) in (
        (top,    (x, y, w, top["w"])),
        (bottom, (x, y + h - bottom["w"], w, bottom["w"])),
        (left,   (x, y, left["w"], h)),
        (right,  (x + w - right["w"], y, right["w"], h)),
    ):
        if b["w"] <= 0.5:
            continue
        bar = _rect_shape(shapes, bx, by, bw, bh)
        set_solid_fill(bar, b["c"]["c"], b["c"]["a"] * op)
        set_line(bar)
        _strip_text_frame(bar)


def _align(value: str):
    return {"left": PP_ALIGN.LEFT, "start": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}.get(value, PP_ALIGN.LEFT)


def _set_line_spacing(paragraph, ratio: float) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("a:lnSpc")):
        pPr.remove(old)
    ln = etree.Element(qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPct"), val=str(int(ratio * 100000)))
    pPr.insert(0, ln)


def add_text_atom(shapes, a: dict) -> None:
    pad_t, pad_r, pad_b, pad_l = a["pad"]
    single_line = a["h"] - pad_t - pad_b < a["lineHeight"] * 1.5
    # Holgura horizontal: las métricas de PowerPoint difieren ligeramente de las del navegador.
    slack = 6 if single_line else 2
    x = a["x"]
    if a["align"] in ("center",):
        x -= slack / 2
    elif a["align"] in ("right", "end"):
        x -= slack
    tb = shapes.add_textbox(px(x), px(a["y"]), px(a["w"] + slack), px(a["h"]))
    if a.get("rot"):
        tb.rotation = a["rot"]
    tf = tb.text_frame
    tf.word_wrap = not single_line
    tf.auto_size = None
    tf.margin_left, tf.margin_right = Pt(pad_l), Pt(pad_r)
    tf.margin_top, tf.margin_bottom = Pt(pad_t), Pt(pad_b)
    # Contenido centrado verticalmente en su caja (flex align-items:center).
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if a.get("alignItems") == "center" else MSO_ANCHOR.TOP

    align = _align(a["align"])
    ratio = a["lineHeight"] / a["fontSize"] if a["fontSize"] else 1.2
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    _set_line_spacing(paragraph, ratio)

    def emit(p, text, run):
        upper = run.get("upper") if run.get("upper") is not None else a["uppercase"]
        r = p.add_run()
        r.text = text.upper() if upper else text
        f = r.font
        f.name = FONT
        f.size = Pt(run.get("size") or a["fontSize"])
        f.bold = bool(run.get("bold"))
        f.italic = bool(run.get("italic"))
        f.color.rgb = RGBColor(*(run.get("color") or a["color"]))
        spacing = run.get("spacing") if run.get("spacing") is not None else a["letterSpacing"]
        if spacing:
            r._r.get_or_add_rPr().set("spc", str(int(round(spacing * 100))))

    # Colapsa espacios en blanco como el navegador (white-space: normal).
    pending_space = False
    first = True
    for run in a["runs"]:
        if run["text"] == "\n":
            paragraph = tf.add_paragraph()
            paragraph.alignment = align
            _set_line_spacing(paragraph, ratio)
            pending_space, first = False, True
            continue
        text = " ".join(run["text"].split())
        lead = run["text"][:1].isspace()
        trail = run["text"][-1:].isspace()
        if not text:
            pending_space = pending_space or (not first)
            continue
        if (pending_space or lead) and not first:
            text = " " + text
        emit(paragraph, text, run)
        first = False
        pending_space = trail


def add_icon_atom(shapes, a: dict, icon_png: bytes) -> None:
    shapes.add_picture(BytesIO(icon_png), px(a["x"]), px(a["y"]), px(a["w"]), px(a["h"]))


# ─── animaciones y transiciones (PresentationML) ──────────────────────────

class _Ids:
    def __init__(self) -> None:
        self.n = 2

    def next(self) -> int:
        self.n += 1
        return self.n


def _float_in_effect(ids: _Ids, spid: int, node_type: str) -> str:
    """Entrada «Flotar hacia dentro» (presetID 42): fundido + subida, como .phase."""
    d = PHASE_DURATION_MS
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (
        f'<p:par><p:cTn id="{ids.next()}" presetID="42" presetClass="entr" presetSubtype="0" '
        f'fill="hold" grpId="0" nodeType="{node_type}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:set><p:cBhvr><p:cTn id="{ids.next()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
        f'{tgt}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
        f'<p:to><p:strVal val="visible"/></p:to></p:set>'
        f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{ids.next()}" dur="{d}"/>{tgt}</p:cBhvr></p:animEffect>'
        f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{ids.next()}" dur="{d}" decel="100000" fill="hold"/>{tgt}'
        f'<p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
        f'<p:tav tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav>'
        f'<p:tav tm="100000"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim>'
        f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{ids.next()}" dur="{d}" decel="100000" fill="hold"/>{tgt}'
        f'<p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
        f'<p:tav tm="0"><p:val><p:strVal val="#ppt_y+{PHASE_RISE:.4f}"/></p:val></p:tav>'
        f'<p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'
        f'</p:childTnLst></p:cTn></p:par>'
    )


def add_phase_timing(slide, steps: list[list[int]]) -> None:
    """steps: una lista de spids por clic (todos los del mismo data-phase aparecen a la vez)."""
    ids = _Ids()
    clicks = []
    for spids in steps:
        outer, inner = ids.next(), ids.next()
        effects = "".join(_float_in_effect(ids, spid, "clickEffect" if k == 0 else "withEffect")
                          for k, spid in enumerate(spids))
        clicks.append(
            f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{effects}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
        )
    bld = "".join(f'<p:bldP spid="{spid}" grpId="0"/>' for spids in steps for spid in spids)
    xml = (
        f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par>'
        f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'{"".join(clicks)}</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'<p:bldLst>{bld}</p:bldLst></p:timing>'
    )
    slide._element.append(etree.fromstring(xml))


def add_push_transition(slide) -> None:
    """«Empujar» desde la derecha (dir=l: el contenido se desplaza hacia la izquierda)."""
    xml = (
        f'<mc:AlternateContent xmlns:mc="{NS_MC}" xmlns:p="{NS_P}">'
        f'<mc:Choice xmlns:p14="{NS_P14}" Requires="p14">'
        f'<p:transition spd="slow" p14:dur="{SLIDE_TRANSITION_MS}"><p:push dir="l"/></p:transition></mc:Choice>'
        f'<mc:Fallback><p:transition spd="slow"><p:push dir="l"/></p:transition></mc:Fallback>'
        f'</mc:AlternateContent>'
    )
    # Orden CT_Slide: cSld, clrMapOvr, transition, timing.
    slide._element.find(qn("p:clrMapOvr")).addnext(etree.fromstring(xml))


def set_background(slide, c) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*c)


# ─── pipeline ─────────────────────────────────────────────────────────────

def build() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"{SOURCE.name} not found: run build_standalone.py first")
    html = SOURCE.read_text(encoding="utf-8")
    # El bloque de Font Awesome es el <style> que contiene "Font Awesome 6 Free".
    fa_start = html.rfind("<style>", 0, html.index('"Font Awesome 6 Free"'))
    fa_css = html[fa_start: html.index("</style>", fa_start) + len("</style>")]

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W_CSS), px(H_CSS)
    blank = prs.slide_layouts[6]
    icon_cache: dict[tuple, bytes] = {}
    report = []

    with sync_playwright() as pw:
        browser = pw.chromium.launch(channel="chrome")
        ctx = browser.new_context(viewport={"width": W_CSS, "height": H_CSS}, offline=True)
        page = ctx.new_page()
        page.goto(SOURCE.as_uri())
        page.evaluate("document.fonts.ready")
        # Sin escalado del .stage: cada slide se mide a 1280×720 nominal.
        page.add_style_tag(content=".stage{transform:none!important}")
        pseudo_n = page.evaluate(JS_PREPARE)
        page.wait_for_timeout(300)
        n = page.locator(".slide-container").count()
        log("source", slides=n, pseudo_elements=pseudo_n)

        # Página aparte, sin fondo y a 4×, para rasterizar cada icono con su color y tamaño.
        hi = browser.new_context(viewport={"width": 400, "height": 400}, device_scale_factor=ICON_SCALE)
        icon_page = hi.new_page()
        icon_page.set_content(ICON_PAGE.format(fa_css=fa_css))
        icon_page.evaluate("document.fonts.ready")

        def icon_png(a: dict) -> bytes:
            key = (a["cls"], a["color"], a["fontSize"], round(a["w"], 1), round(a["h"], 1))
            if key not in icon_cache:
                icon_page.evaluate(
                    "([cls, color, size, w, h]) => { const i = document.getElementById('icon');"
                    " i.className = cls; i.style.color = color; i.style.fontSize = size;"
                    " i.style.width = w + 'px'; i.style.height = h + 'px'; i.style.textAlign = 'center'; }",
                    [a["cls"], a["color"], a["fontSize"], a["w"], a["h"]])
                icon_cache[key] = icon_page.locator("#icon").screenshot(omit_background=True, scale="device")
            return icon_cache[key]

        for i in range(n):
            data = page.evaluate(JS_EXTRACT, i)
            slide = prs.slides.add_slide(blank)
            set_background(slide, data["bg"])
            groups: dict[int, object] = {}
            counts = {"shape": 0, "text": 0, "icon": 0}
            for a in data["atoms"]:
                target = slide.shapes
                if a.get("phase"):
                    key = a["phase"]["key"]
                    if key not in groups:
                        groups[key] = slide.shapes.add_group_shape()
                    target = groups[key].shapes
                if a["kind"] == "shape":
                    add_shape_atom(target, a)
                elif a["kind"] == "text":
                    add_text_atom(target, a)
                elif a["kind"] == "icon":
                    add_icon_atom(target, a, icon_png(a))
                counts[a["kind"]] += 1

            animated = 0
            if groups:
                by_num: dict[int, list[int]] = {}
                for key, group in sorted(groups.items()):
                    group.name = f"Fase {data['phases'][key]}"
                    by_num.setdefault(data["phases"][key], []).append(group.shape_id)
                add_phase_timing(slide, [by_num[k] for k in sorted(by_num)])
                animated = len(groups)
            if i > 0:
                add_push_transition(slide)
            report.append({"slide": i + 1, **counts, "animated_groups": animated})
            log("slide", slide=i + 1, **counts, animated_groups=animated)

        browser.close()

    prs.save(OUT_PPTX)
    log("done", output=OUT_PPTX.name, kb=OUT_PPTX.stat().st_size // 1024, icons=len(icon_cache))


if __name__ == "__main__":
    build()
